"""
Campaign service layer.

Contains:
- Document upload validation (extension, MIME, zip safety)
- S3 staging
- Redis batch tracker CRUD
- Generic Document Text Extraction (PDF, DOCX, TXT)
- Generic LLM Document Fields Extraction
- Generic LLM Screening
"""

from __future__ import annotations

import asyncio
import io
import os
import re
import unicodedata
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from fastapi import HTTPException, UploadFile, status
from redis.asyncio import Redis

from app.core.config import settings
from app.core.s3 import upload_fileobj_to_s3
from app.integrations.ai.factory import StructuredExtractionProviderFactory
from app.integrations.ai.providers.base import StructuredExtractionProvider


# =========================================================================
# DOCUMENT UPLOAD VALIDATION
# =========================================================================

ALLOWED_EXTENSIONS: set[str] = {".pdf", ".docx", ".txt", ".zip"}
BLOCKED_EXTENSIONS: set[str] = {".exe", ".js", ".html", ".csv", ".bat", ".sh", ".cmd", ".msi", ".dll"}

EXT_TO_MIME: dict[str, set[str]] = {
    ".pdf": {"application/pdf"},
    ".docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/zip",
    },
    ".txt": {"text/plain", "application/octet-stream"},
    ".zip": {"application/zip", "application/x-zip-compressed", "application/octet-stream"},
}

_ZIP_MAX_UNCOMPRESSED_BYTES = 500 * 1024 * 1024  # 500 MB
_ZIP_MAX_RATIO = 50


def validate_document_uploads(files: list[UploadFile]) -> list[UploadFile]:
    if not files:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="At least one file is required.",
        )

    for f in files:
        if not f.filename:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Every uploaded file must have a filename.",
            )
        ext = Path(f.filename).suffix.lower()

        if ext in BLOCKED_EXTENSIONS:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Blocked file type: {f.filename}",
            )
        if ext not in ALLOWED_EXTENSIONS:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Unsupported file type: {f.filename}. Allowed: pdf, docx, txt, zip",
            )

        expected = EXT_TO_MIME.get(ext, set())
        actual = f.content_type or ""
        if actual and expected and actual not in expected:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"MIME mismatch for {f.filename}: got {actual}",
            )

    return files


async def validate_zip_safety(data: bytes) -> list[str]:
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Uploaded zip file is corrupt or not a valid zip.",
        )

    infos = zf.infolist()
    max_count = getattr(settings, "RESUME_MAX_FILE_COUNT", 500)
    if len(infos) > max_count:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Zip contains {len(infos)} entries (max {max_count}).",
        )

    total_uncompressed = sum(i.file_size for i in infos)
    compressed_size = len(data)
    if total_uncompressed > _ZIP_MAX_UNCOMPRESSED_BYTES:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Zip file exceeds maximum uncompressed size.",
        )
    if compressed_size > 0 and (total_uncompressed / compressed_size) > _ZIP_MAX_RATIO:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Zip compression ratio is suspiciously high.",
        )

    safe_names: list[str] = []
    for info in infos:
        name = info.filename
        if info.is_dir():
            continue
        if (info.external_attr >> 16) & 0o120000 == 0o120000:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Symlink found: {name}")
        if name.startswith("/") or name.startswith("\\") or ".." in name:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Path traversal: {name}")

        basename = Path(name).name
        if basename.startswith(".") or basename in {"__MACOSX", "Thumbs.db", "desktop.ini"} or "__MACOSX" in name:
            continue

        if Path(name).suffix.lower() == ".zip":
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Nested zip found: {name}")

        ext = Path(name).suffix.lower()
        if ext not in {".pdf", ".docx", ".txt"}:
            continue

        safe_names.append(name)

    if not safe_names:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Zip contains no supported files.",
        )

    return safe_names


# =========================================================================
# S3 STAGING
# =========================================================================

async def stage_files_to_s3(
    batch_id: str,
    files: list[UploadFile],
) -> tuple[str, str]:
    s3_prefix = f"document-batches/{batch_id}/original"
    source_type = "single" if len(files) == 1 else "multi"

    for f in files:
        if Path(f.filename).suffix.lower() == ".zip":
            source_type = "zip"

        key = f"{s3_prefix}/{f.filename}"
        content_type = f.content_type or "application/octet-stream"
        await asyncio.to_thread(
            upload_fileobj_to_s3, key, f.file, content_type=content_type
        )

    return s3_prefix, source_type


# =========================================================================
# REDIS BATCH TRACKER
# =========================================================================

async def create_batch_tracker(redis: Redis, batch_id: str, file_count: int) -> None:
    now = datetime.now(timezone.utc).isoformat()
    await redis.hset(  # type: ignore
        f"job:{batch_id}",
        mapping={
            "status": "QUEUED",
            "total_files": str(file_count),
            "processed": "0",
            "failed": "0",
            "created_at": now,
            "updated_at": now,
        },
    )
    await redis.expire(f"job:{batch_id}", 7 * 86400)


async def update_batch_progress(
    redis: Redis,
    batch_id: str,
    *,
    total_files: int | None = None,
    processed_incr: int = 0,
    failed_incr: int = 0,
    status: str | None = None,
) -> None:
    pipe = redis.pipeline(transaction=True)
    if total_files is not None:
        pipe.hset(f"job:{batch_id}", "total_files", str(total_files))
    if processed_incr:
        pipe.hincrby(f"job:{batch_id}", "processed", processed_incr)
    if failed_incr:
        pipe.hincrby(f"job:{batch_id}", "failed", failed_incr)
    if status:
        pipe.hset(f"job:{batch_id}", "status", status)
    pipe.hset(f"job:{batch_id}", "updated_at", datetime.now(timezone.utc).isoformat())
    await pipe.execute()


async def complete_batch(redis: Redis, batch_id: str) -> None:
    now = datetime.now(timezone.utc).isoformat()
    await redis.hset(  # type: ignore
        f"job:{batch_id}",
        mapping={
            "status": "COMPLETED",
            "finished_at": now,
            "updated_at": now,
        },
    )


async def fail_batch(redis: Redis, batch_id: str, reason: str = "") -> None:
    now = datetime.now(timezone.utc).isoformat()
    await redis.hset(  # type: ignore
        f"job:{batch_id}",
        mapping={
            "status": "FAILED",
            "finished_at": now,
            "updated_at": now,
            "error": reason[:500],
        },
    )


async def get_batch_status(redis: Redis, batch_id: str) -> dict[str, str] | None:
    data = await redis.hgetall(f"job:{batch_id}")  # type: ignore
    return data or None


# =========================================================================
# ARQ JOB ENQUEUE
# =========================================================================

async def enqueue_document_upload_batch(
    redis: Redis,
    *,
    batch_id: str,
    campaign_id: str,
    s3_prefix: str,
    source_type: str,
) -> None:
    from arq import create_pool
    from arq.connections import RedisSettings

    pool = await create_pool(RedisSettings.from_dsn(settings.REDIS_URL))
    await pool.enqueue_job(
        "process_document_upload_batch",
        batch_id=batch_id,
        campaign_id=campaign_id,
        s3_prefix=s3_prefix,
        source_type=source_type,
    )
    await pool.aclose()


async def enqueue_campaign_screening(
    redis: Redis,
    *,
    batch_id: str,
    campaign_id: str,
    candidate_ids: list[str] | None = None,
) -> None:
    from arq import create_pool
    from arq.connections import RedisSettings

    pool = await create_pool(RedisSettings.from_dsn(settings.REDIS_URL))
    await pool.enqueue_job(
        "screen_campaign_candidates",
        batch_id=batch_id,
        campaign_id=campaign_id,
        candidate_ids=candidate_ids,
    )
    await pool.aclose()


# =========================================================================
# TEXT EXTRACTION
# =========================================================================

def normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = re.sub(r"[\x00-\x08\x0b-\x1f\x7f]", "", text)
    text = re.sub(r"\r\n?", "\n", text)
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

def extract_text_from_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return normalize_text("\n".join(pages))


def extract_text_from_docx(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts: list[str] = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return normalize_text("\n".join(parts))


def extract_text_from_txt(data: bytes) -> str:
    for enc in ("utf-8", "latin-1", "cp1252"):
        try:
            return normalize_text(data.decode(enc))
        except (UnicodeDecodeError, ValueError):
            continue
    return normalize_text(data.decode("utf-8", errors="replace"))


def extract_document_text(data: bytes, ext: str) -> str | None:
    try:
        if ext == ".pdf":
            return extract_text_from_pdf(data)
        if ext == ".docx":
            return extract_text_from_docx(data)
        if ext == ".txt":
            return extract_text_from_txt(data)
    except Exception:
        return None
    return None

def validate_magic_bytes(data: bytes, ext: str) -> bool:
    if ext == ".pdf":
        return data[:5] == b"%PDF-"
    if ext == ".docx":
        return data[:4] == b"PK\x03\x04"
    if ext == ".txt":
        sample = data[:512]
        control = sum(1 for b in sample if b < 0x09 or (0x0E <= b <= 0x1F) or b == 0x7F)
        return control <= 2
    return False

# =========================================================================
# GENERIC LLM EXTRACTION & SCREENING
# =========================================================================

async def extract_document_fields_llm(
    document_text: str,
    existing_fields: dict | None = None,
    extraction_provider: StructuredExtractionProvider | None = None,
) -> dict[str, Any]:
    if not document_text or not document_text.strip():
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Document text cannot be empty for structured extraction.",
        )
    if extraction_provider is None:
        extraction_provider = StructuredExtractionProviderFactory.build()

    return await extraction_provider.extract(document_text, existing_fields)


async def screen_document_llm(
    candidate_fields: dict,
    campaign_fields: dict,
    extraction_provider: StructuredExtractionProvider | None = None,
) -> dict[str, Any]:
    if extraction_provider is None:
        extraction_provider = StructuredExtractionProviderFactory.build()

    return await extraction_provider.screen_candidate(
        candidate_text="",
        candidate_fields=candidate_fields or {},
        campaign_text="",
        campaign_fields=campaign_fields or {},
    )

# =========================================================================
# WEBHOOK PROCESSING
# =========================================================================

from app.domains.campaigns.schemas import CallWebhookPayload
from app.domains.campaigns.models import CallScreening, Candidate, Campaign
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

async def process_call_webhook(db: AsyncSession, payload: CallWebhookPayload) -> None:
    # 1. Fetch Candidate and Campaign
    result = await db.execute(
        select(Candidate).where(Candidate.id == payload.candidate_id)
    )
    candidate = result.scalar_one_or_none()
    if not candidate:
        raise HTTPException(status_code=404, detail="Candidate not found")
        
    result = await db.execute(
        select(Campaign).where(Campaign.id == payload.campaign_id)
    )
    campaign = result.scalar_one_or_none()
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found")

    # 2. Extract Data from Transcript & Screen using Generic Engine
    extraction_provider = StructuredExtractionProviderFactory.build()
    
    transcript = payload.transcript or ""
    
    # We pass existing fields so the LLM has context, but we instruct it 
    # to extract new fields or updated info.
    extracted_fields = await extraction_provider.extract(transcript, candidate.extracted_fields)
    
    # Screen Candidate based on the new transcript + fields vs campaign requirements
    screening_result = await extraction_provider.screen_candidate(
        candidate_text=transcript,
        candidate_fields=extracted_fields,
        campaign_text=campaign.raw_text or "",
        campaign_fields=campaign.required_fields or {},
    )
    
    # 3. Create CallScreening record
    call_screening = CallScreening(
        campaign_id=campaign.id,
        candidate_id=candidate.id,
        transcript=transcript,
        recording_url=payload.recording_url,
        match_score=screening_result.get("match_score", 0.0),
        matched_fields=screening_result.get("matched_fields", {}),
        unmatched_fields=screening_result.get("unmatched_fields", {}),
        summary=screening_result.get("summary", ""),
    )
    db.add(call_screening)
    
    # 4. Update Candidate Fields
    # The rule is: data extracted from the direct phone call overwrites existing fields.
    merged_fields = dict(candidate.extracted_fields or {})
    merged_fields.update(extracted_fields)
    candidate.extracted_fields = merged_fields
    
    # If the LLM returned primary fields, update them
    if "name" in extracted_fields and extracted_fields["name"]:
        candidate.name = extracted_fields["name"]
    if "email" in extracted_fields and extracted_fields["email"]:
        candidate.email = extracted_fields["email"]
    if "phone" in extracted_fields and extracted_fields["phone"]:
        candidate.phone = extracted_fields["phone"]
        
    await db.commit()
    
    # 5. Notify Orchestrator to proceed to next step
    from app.domains.campaigns.orchestrator import on_step_completed
    await on_step_completed(
        db=db,
        candidate_id=candidate.id,
        service_name="outbound_call",
        payload={
            "call_id": payload.call_id,
            "status": payload.status,
            "match_score": call_screening.match_score,
        }
    )